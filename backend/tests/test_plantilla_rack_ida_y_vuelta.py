# -*- coding: utf-8 -*-
"""La plantilla de Rooms: baja, se llena, sube.

Owner, 2026-10-16: *«por qué no creás un download file para actualizar los
datos vía upload»*.

Lo que se prueba acá es el **ida y vuelta**: lo que baja tiene que poder subir.
Un archivo que se genera lindo y no se puede leer de vuelta es el modo de falla
de toda plantilla — y no se descubre hasta que alguien pasó una hora
llenándola.
"""
from decimal import Decimal

import pytest

from app.export.rack_rooms_xlsx import (clave_de, construir_plantilla,
                                        leer_plantilla, partir_clave)

SEGMENTOS = [
    {"code": "AGENCIAS", "nombre": "Agencias", "propia": False,
     # La comisión sube a mitad de año: es justo el caso que motivó el cambio.
     "comisiones": [0.25] * 6 + [0.28] * 6,
     "noches": [25, 65, 45, 25, 20, 20, 50, 30, 20, 20, 20, 60],
     "tarifas": [0] * 12},
    {"code": "LONG_TERM", "nombre": "Long term rentals", "propia": True,
     "comisiones": [0.0] * 12,
     "noches": [30, 15, 0, 30, 30, 30, 30, 30, 30, 30, 30, 30],
     "tarifas": [75, 100, 0, 75, 75, 75, 75, 75, 75, 75, 75, 75]},
    {"code": "CPL", "nombre": "CPL | Uso casa | Canjes", "propia": False,
     "comisiones": [0.0] * 12, "noches": [0] * 12, "tarifas": [0] * 12},
]


@pytest.fixture
def archivo():
    return construir_plantilla(SEGMENTOS, "Ojochal Gardens", "BUDGET Working",
                               2027)


# ── El ida y vuelta ─────────────────────────────────────────────────────────
def test_lo_que_baja_se_puede_leer_de_vuelta(archivo):
    filas = leer_plantilla(archivo)
    claves = [c for c, _, _ in filas]
    assert "AGENCIAS" in claves
    assert "LONG_TERM" in claves
    assert "LONG_TERM|TARIFA" in claves


def test_los_numeros_vuelven_IGUALES(archivo):
    """Bajar y subir sin tocar nada no puede cambiar un solo número."""
    porclave = {c: (com, meses) for c, com, meses in leer_plantilla(archivo)}
    _, meses = porclave["AGENCIAS"]
    assert [int(x) for x in meses] == SEGMENTOS[0]["noches"]
    _, tarifas = porclave["LONG_TERM|TARIFA"]
    assert [float(x) for x in tarifas] == [float(x) for x in SEGMENTOS[1]["tarifas"]]


def test_la_comision_vuelve_como_fraccion_y_POR_MES(archivo):
    """0,25 y no 25: el formato de Excel muestra «25 %» pero el valor es 0,25.

    Si volviera como 25, el endpoint lo rechazaría por estar fuera de rango —
    que es mejor que aplicarlo, pero igual sería una plantilla rota.

    Y vuelve mes a mes: la comisión que sube a mitad de año tiene que llegar
    como seis meses de 0,25 y seis de 0,28, no como un promedio.
    """
    porclave = {c: meses for c, _, meses in leer_plantilla(archivo)}
    com = [float(x) for x in porclave["AGENCIAS|COMISION"]]
    assert com == [0.25] * 6 + [0.28] * 6


def test_un_segmento_de_tarifa_propia_NO_trae_fila_de_comision(archivo):
    """Su tarifa YA es la neta: traer la fila invitaría a llenarla, y el
    importador la rechazaría después de que alguien la escribió."""
    claves = [c for c, _, _ in leer_plantilla(archivo)]
    assert "LONG_TERM|COMISION" not in claves
    assert "AGENCIAS|COMISION" in claves


def test_un_segmento_sin_tarifa_propia_NO_trae_fila_de_tarifa(archivo):
    """Si la trajera, se podría escribir una tarifa que el motor ignora — y el
    que la escribió esperaría verla en el ingreso."""
    claves = [c for c, _, _ in leer_plantilla(archivo)]
    assert "AGENCIAS|TARIFA" not in claves


# ── La clave ────────────────────────────────────────────────────────────────
def test_la_clave_lleva_el_tipo_adentro():
    """En una columna aparte, un `ordenar` de Excel dejaría la fila apareada
    contra otra cosa sin que nadie lo note."""
    assert clave_de("OTA", "NOCHES") == "OTA"
    assert clave_de("LONG_TERM", "TARIFA") == "LONG_TERM|TARIFA"
    assert partir_clave("OTA") == ("OTA", "NOCHES")
    assert partir_clave("LONG_TERM|TARIFA") == ("LONG_TERM", "TARIFA")
    assert clave_de("OTA", "COMISION") == "OTA|COMISION"
    assert partir_clave("OTA|COMISION") == ("OTA", "COMISION")


def test_la_clave_tolera_espacios():
    """Excel y el copiar/pegar meten espacios donde no había."""
    assert partir_clave("  OTA  ") == ("OTA", "NOCHES")
    assert partir_clave("LONG_TERM | tarifa ") == ("LONG_TERM", "TARIFA")


# ── El encabezado ───────────────────────────────────────────────────────────
def test_el_encabezado_se_busca_por_contenido(archivo):
    """El día que la leyenda crezca una línea, el archivo que alguien tenía a
    medio llenar tiene que seguir subiendo."""
    import inspect

    import app.export.rack_rooms_xlsx as mod
    fuente = inspect.getsource(mod.leer_plantilla)
    assert '"clave" in vals' in fuente
    assert "FILA_ENCABEZADO" not in fuente


def test_un_archivo_sin_encabezado_falla_claro():
    from openpyxl import Workbook

    from app.export.excel_base import workbook_to_bytes
    wb = Workbook()
    wb.active["A1"] = "cualquier cosa"
    with pytest.raises(ValueError):
        leer_plantilla(workbook_to_bytes(wb))


def test_los_meses_se_ubican_por_rotulo(archivo):
    """Una columna metida en el medio no puede correr los doce meses."""
    import inspect

    import app.export.rack_rooms_xlsx as mod
    assert "cab.index(m.lower())" in inspect.getsource(mod.leer_plantilla)


# ── Lo que NO se puede editar ───────────────────────────────────────────────
def test_la_hoja_va_protegida(archivo):
    """El archivo vuelve y se aparea por `clave`: una columna corrida a mano
    rompe el apareo y el que sube no se entera."""
    import io

    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(archivo)).active
    assert ws.protection.sheet is True


def test_la_clave_queda_bloqueada(archivo):
    import io

    from openpyxl import load_workbook
    from app.export.rack_rooms_xlsx import COL_CLAVE, FILA_ENCABEZADO
    ws = load_workbook(io.BytesIO(archivo)).active
    assert ws.cell(FILA_ENCABEZADO + 1, COL_CLAVE).protection.locked is True


def test_las_noches_quedan_desbloqueadas(archivo):
    """Es lo único que se llena: si queda bloqueado, la plantilla no sirve."""
    import io

    from openpyxl import load_workbook
    from app.export.rack_rooms_xlsx import COL_ENE, FILA_ENCABEZADO
    ws = load_workbook(io.BytesIO(archivo)).active
    assert ws.cell(FILA_ENCABEZADO + 1, COL_ENE).protection.locked is False


def test_el_total_es_FORMULA_y_no_numero(archivo):
    """Si fuera número, cambiar un mes dejaría el total contradiciendo a la
    fila de al lado y no se sabría cuál mirar."""
    import io

    from openpyxl import load_workbook
    from app.export.rack_rooms_xlsx import COL_TOTAL, FILA_ENCABEZADO
    ws = load_workbook(io.BytesIO(archivo)).active   # sin data_only: la fórmula
    v = ws.cell(FILA_ENCABEZADO + 1, COL_TOTAL).value
    assert isinstance(v, str) and v.startswith("=SUM(")


# ── El endpoint que la lee ──────────────────────────────────────────────────
def test_el_upload_es_en_seco_por_omision():
    import inspect

    from app.api import rack_rooms_api as api
    firma = inspect.signature(api.subir_plantilla)
    assert firma.parameters["aplicar"].default.default is False


def test_una_celda_vacia_NO_es_un_cero():
    """El descuido más común —borrar una celda, filtrar y no darse cuenta— no
    puede poner un mes en cero con el total cuadrando consigo mismo."""
    import inspect

    from app.api import rack_rooms_api as api
    fuente = inspect.getsource(api.subir_plantilla)
    assert "if valor is None:" in fuente
    assert "continue" in fuente


def test_el_upload_en_seco_no_deja_nada_escrito():
    """Se tocan objetos de la sesión al comparar. Sin `rollback` ese cambio
    viaja en el `commit` de cualquier otra ruta."""
    import inspect

    from app.api import rack_rooms_api as api
    fuente = inspect.getsource(api.subir_plantilla)
    assert "await db.rollback()" in fuente
    assert fuente.index("await db.rollback()") < fuente.index("await db.commit()")


def test_el_upload_rederiva_al_aplicar():
    """Cambiar noches o comisión cambia la tarifa y el ingreso. Sin rederivar,
    la tabla estándar —la que lee el consolidado— queda con lo de antes."""
    import inspect

    from app.api import rack_rooms_api as api
    assert "_rederivar(" in inspect.getsource(api.subir_plantilla)


def test_el_upload_respeta_el_candado():
    import inspect

    from app.api import rack_rooms_api as api
    assert "await candado(db, scenario_id)" in inspect.getsource(api.subir_plantilla)
