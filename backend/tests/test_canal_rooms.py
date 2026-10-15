# -*- coding: utf-8 -*-
"""Rooms por canal: el modelo alternativo para quien no tiene historia por villa.

Owner, 2026-10-15: *«no tenemos historial de ventas totales, ocupación y tarifa
porque no hay históricos por tipo de habitación como sí tengo en FinPlan CWL…
Sales and Marketing me proponen hacerlo de otra manera, por canal»*.

Lo que se cuida acá son las tres cosas que, si se rompen, **cuadran solas**:

1. **El triángulo noches / tarifa / ingreso.** Cuál de los tres manda cambia
   según la versión: en la proyección se digita la tarifa y el ingreso se
   deriva; en el histórico se digita el ingreso y la tarifa es el promedio que
   salió del PMS. Derivar al revés mueve plata sin avisar.

2. **El CPL fuera del ADR.** Las noches de cortesía ocupan una villa pero no
   generan ingreso. Si entran al denominador, el ADR se diluye y deja de ser
   comparable contra el PMS — y el total de ingreso no cambia, así que nada lo
   señala.

3. **El interruptor.** Sin `rooms_source = 'canales'` el modelo no toca el P&L.
   CWL tiene buena historia por tipo de habitación y se queda con el suyo.
"""
from decimal import Decimal as D

import pytest

from app.models.canal_rooms import (SEGMENTOS_SEMILLA, CanalRoomsEntry,
                                    calcular_mes, derivar)


def _e(code, noches, tarifa=0, ingreso=0):
    return CanalRoomsEntry(segmento_code=code, month=1, noches=D(str(noches)),
                           tarifa=D(str(tarifa)), ingreso=D(str(ingreso)))


GENERA = {c: g for c, _n, g in SEGMENTOS_SEMILLA}


# ── El triangulo ─────────────────────────────────────────────────────────────

def test_con_tarifa_digitada_el_ingreso_se_deriva():
    """Es la proyección: se digitan noches y tarifa."""
    assert derivar(100, "592.62", None, True) == (D("592.62"), D("59262.00"))


def test_con_ingreso_digitado_la_tarifa_es_el_promedio():
    """Es el histórico: el PMS da noches e ingreso, la tarifa sale de dividir."""
    t, i = derivar(86, None, "50965", False)
    assert t == D("592.62")
    assert i == D("50965")


def test_sin_noches_no_se_divide_por_nada():
    assert derivar(0, "500", None, True) == (D("0"), D("0"))
    assert derivar(0, None, "1000", False)[0] == D("0")


def test_el_caso_del_libro_de_sales():
    """Enero de Agencias, tal como está en el archivo.

    OTB: 86 noches y $50.965 → tarifa $592,62. La proyección toma esa tarifa
    para 100 noches → $59.261,63, que es el `D9` del original.
    """
    tarifa, _ = derivar(86, None, "50965", False)
    _, ingreso = derivar(100, tarifa, None, True)
    assert abs(ingreso - D("59261.63")) < D("0.60")


# ── El CPL ───────────────────────────────────────────────────────────────────

def test_el_cpl_ocupa_la_villa_pero_no_entra_al_adr():
    """La regla que hace que el cuadro cierre contra el PMS."""
    r = calcular_mes([_e("AGENCIAS", 100, ingreso=50000),
                      _e("CPL", 10, ingreso=0)], GENERA, 403)
    assert r["noches_ingreso"] == D("100")      # el ADR divide por esto
    assert r["noches_totales"] == D("110")      # la villa estuvo ocupada
    assert r["adr"] == D("500")                 # NO 454,55
    assert r["ocupacion"] == D("100") / D("403")
    assert r["ocupacion_total"] == D("110") / D("403")


def test_sin_cpl_las_dos_ocupaciones_coinciden():
    r = calcular_mes([_e("AGENCIAS", 100, ingreso=50000)], GENERA, 403)
    assert r["noches_ingreso"] == r["noches_totales"]
    assert r["ocupacion"] == r["ocupacion_total"]


def test_el_revpar_divide_por_las_disponibles():
    r = calcular_mes([_e("AGENCIAS", 100, ingreso=50000)], GENERA, 403)
    assert r["revpar"] == D("50000") / D("403")


def test_mes_vacio_no_revienta():
    r = calcular_mes([], GENERA, 403)
    assert r["ingreso"] == D("0") and r["adr"] == D("0") and r["revpar"] == D("0")


def test_sin_disponibles_la_ocupacion_es_cero_y_no_infinito():
    r = calcular_mes([_e("AGENCIAS", 10, ingreso=5000)], GENERA, 0)
    assert r["ocupacion"] == D("0")
    assert r["adr"] == D("500")          # el ADR no depende de las disponibles


# ── El catalogo ──────────────────────────────────────────────────────────────

def test_los_ocho_segmentos_del_libro_de_sales():
    assert len(SEGMENTOS_SEMILLA) == 8
    codes = [c for c, _n, _g in SEGMENTOS_SEMILLA]
    assert "CPL" in codes


def test_solo_el_cpl_no_genera_ingreso():
    """Si alguien marca otro como cortesía, el ADR de ese mes salta sin que el
    ingreso se mueva — y eso no se nota mirando el total."""
    sin = [c for c, _n, g in SEGMENTOS_SEMILLA if not g]
    assert sin == ["CPL"]


# ── El importador ───────────────────────────────────────────────────────────
#
# ⚠️ El libro se arma ACA, no se lee del disco. Una prueba atada a un archivo
# de una maquina pasa en todas las demas sin mirar nada — que es el peor
# resultado posible. Los numeros son los de enero del libro real de Sales
# (`Presupuesto_Oxygen_2027_Modelo_v2`), asi que sigue cuadrando contra el.

#: Enero del libro real: {code: (noches, ingreso)} de la columna «Real 2026».
ENERO_REAL_2026 = {
    "Agencias": (80, 41219), "Direct (website / Rsv)": (121, 68539),
    "Costa Rica Collection direct": (79, 16876), "OTA's": (44, 34794),
    "CPL | Uso casa | Canjes": (6, 0),
}


def _libro_como_el_de_sales(tmp_path):
    """Un `Resumen Manual` con la MISMA forma que manda Sales.

    Incluye las dos trampas del original: las metricas van indentadas y se
    llaman «Ingreso $» con el simbolo, y despues de los segmentos viene un
    bloque TOTALES que esta en la misma columna y NO es un segmento.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen Manual"
    ws["A1"] = "OXYGEN JUNGLE VILLAS · RESUMEN MANUAL POR SEGMENTO"
    ws["C5"] = "Enero"
    ws["I5"] = "Febrero"
    ws["BW5"] = "TOTAL AÑO"
    ws["A6"] = "Segmento / KPI"
    for base in (3, 9):                      # enero y febrero
        # ⚠️ El salto de línea va DENTRO del rótulo: así lo manda Sales
        # («Proyección 2027⏎10/26»), y por eso el importador normaliza espacios
        # antes de comparar en vez de buscar el texto tal cual.
        for k, v in enumerate(("Budget 2027", "Proyección 2027\n10/26",
                               "OTB 2027\n10/26", "Budget 2026",
                               "Real 2026", "Real 2025")):
            ws.cell(6, base + k, v)
    ws.cell(6, 75, "Var Budget vs 2026")     # del bloque TOTAL AÑO: no es version

    f = 7
    for nombre, (noches, ingreso) in ENERO_REAL_2026.items():
        ws.cell(f, 1, nombre)                # el segmento NO va indentado
        ws.cell(f + 1, 1, "  Noches")        # las metricas SI
        ws.cell(f + 2, 1, "  Ingreso $")
        ws.cell(f + 3, 1, "  Tarifa Promedio")
        ws.cell(f + 4, 1, "  % Ocupación")
        ws.cell(f + 1, 7, noches)            # col G = «Real 2026» de enero
        ws.cell(f + 2, 7, ingreso)
        f += 6

    # El bloque que no hay que leer
    ws.cell(f, 1, "TOTALES")
    for k, et in enumerate(("Habitaciones disponibles",
                            "Noches ocupadas (generan ingreso)",
                            "Total Ingresos $", "Tarifa Promedio (ADR)")):
        ws.cell(f + 1 + k, 1, et)
        ws.cell(f + 1 + k, 7, 99999)         # si se leyera, el total se duplica

    ruta = tmp_path / "sales.xlsx"
    wb.save(ruta)
    return str(ruta)


def _catalogo():
    from app.importers.canal_rooms_importer import _norm
    return {_norm(n): c for c, n, _g in SEGMENTOS_SEMILLA}


def test_el_importador_lee_noches_e_ingreso_por_segmento(tmp_path):
    from app.importers.canal_rooms_importer import leer
    r = leer(_libro_como_el_de_sales(tmp_path), "Real 2026", _catalogo())
    assert r["datos"][("AGENCIAS", 1)]["noches"] == D("80")
    assert r["datos"][("AGENCIAS", 1)]["ingreso"] == D("41219")
    assert r["datos"][("OTA", 1)]["ingreso"] == D("34794")
    # `meses` son los meses donde se ENCONTRÓ la columna de esa versión, no
    # donde hay dato: febrero existe en el libro aunque venga vacío, y decirlo
    # permite distinguir «no lo mandaron» de «lo mandaron en cero».
    assert r["meses"] == [1, 2]


def test_el_importador_no_lee_los_totales_como_segmentos(tmp_path):
    """⚠️ «Total Ingresos $» y «Habitaciones disponibles» estan en la misma
    columna que los segmentos. Leerlos duplicaria el ingreso del mes."""
    from app.importers.canal_rooms_importer import leer
    r = leer(_libro_como_el_de_sales(tmp_path), "Real 2026", _catalogo())
    assert not r["desconocidos"], r["desconocidos"]
    assert {c for c, _m in r["datos"]} <= set(_catalogo().values())
    total = sum(d.get("ingreso", D(0)) for d in r["datos"].values())
    assert total == D(str(sum(i for _n, i in ENERO_REAL_2026.values())))


def test_el_importador_reconoce_la_metrica_con_el_simbolo(tmp_path):
    """«Ingreso $», no «Ingreso». Sin el simbolo el importador devuelve todas
    las noches y CERO ingreso — peor que fallar, porque parece que el archivo
    viniera sin plata."""
    from app.importers.canal_rooms_importer import leer
    r = leer(_libro_como_el_de_sales(tmp_path), "Real 2026", _catalogo())
    assert any(d.get("ingreso") for d in r["datos"].values())


def test_un_segmento_que_no_esta_en_el_catalogo_no_se_inventa(tmp_path):
    from app.importers.canal_rooms_importer import leer
    cat = dict(_catalogo())
    cat.pop("ota's", None)
    r = leer(_libro_como_el_de_sales(tmp_path), "Real 2026", cat)
    assert "OTA's" in r["desconocidos"]
    assert not any(c == "OTA" for c, _m in r["datos"])


def test_las_versiones_no_incluyen_las_del_bloque_total_anio(tmp_path):
    """El bloque TOTAL AÑO trae comparaciones ya calculadas («Var Budget vs
    2026»): no son versiones importables."""
    from app.importers.canal_rooms_importer import versiones_del_libro
    vs = versiones_del_libro(_libro_como_el_de_sales(tmp_path))
    assert "Real 2026" in vs
    assert not any(v.startswith("Var ") for v in vs), vs


def test_el_mes_sale_de_la_banda_y_no_de_la_posicion(tmp_path):
    """Las columnas de febrero son otras; la version se busca por texto bajo su
    banda de mes, no contando columnas."""
    from app.importers.canal_rooms_importer import leer
    from openpyxl import load_workbook
    ruta = _libro_como_el_de_sales(tmp_path)
    wb = load_workbook(ruta)
    wb["Resumen Manual"].cell(8, 13, 55)      # col M = «Real 2026» de febrero
    wb.save(ruta)
    r = leer(ruta, "Real 2026", _catalogo())
    assert r["datos"][("AGENCIAS", 2)]["noches"] == D("55")
    assert sorted(r["meses"]) == [1, 2]


# ── El interruptor ───────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_sin_conectar_el_modelo_no_toca_el_pl():
    """CWL se queda con el modelo de tarifas, y la condición es el escenario."""
    from app.engine.recalculate import aplicar_canales_rooms

    class _Esc:
        id, hotel_id, rooms_source = "S1", "CWL", "tarifas"

    from app.engine.revenue_calculator import RevenueResult
    r = RevenueResult(month=1, year=2027)
    r.rooms = D("999")
    aplico = await aplicar_canales_rooms(None, _Esc(), {1: r}, [])
    assert aplico is False
    assert r.rooms == D("999")
